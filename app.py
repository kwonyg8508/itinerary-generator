import streamlit as st
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import io
import json
import time
from google import genai
from google.genai import types

# 페이지 기본 설정
st.set_page_config(page_title="유럽 랜드사 일정표 생성기", layout="wide", page_icon="🇪🇺")

# 레트로 2D 픽셀 아트 CSS 및 애니메이션 스타일
pixel_art_css = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Press+Start+2P&family=Noto+Sans+KR:wght@400;700&display=swap');

.retro-container {
    background: #1e1e2f;
    border: 4px solid #4a4a75;
    border-radius: 12px;
    padding: 16px;
    box-shadow: 0 8px 0 #11111b;
    margin-bottom: 20px;
    font-family: 'Noto Sans KR', sans-serif;
    color: #f1f1f1;
}

.retro-header {
    display: flex;
    justify-content: space-between;
    align-items: center;
    border-bottom: 2px dashed #4a4a75;
    padding-bottom: 8px;
    margin-bottom: 12px;
}

.pixel-badge {
    background: #ff5e7e;
    color: white;
    font-size: 11px;
    padding: 4px 8px;
    border-radius: 4px;
    font-weight: bold;
    letter-spacing: 0.5px;
}

.retro-stage {
    display: flex;
    align-items: center;
    gap: 20px;
    background: #282846;
    padding: 15px;
    border-radius: 8px;
}

/* 2D 픽셀 캐릭터 박스 */
.character-box {
    width: 80px;
    height: 80px;
    background: #3b3b63;
    border: 3px solid #ffde59;
    border-radius: 8px;
    display: flex;
    align-items: center;
    justify-content: center;
    font-size: 42px;
    position: relative;
    user-select: none;
}

/* 대기 모드: 둥실둥실 숨쉬기 */
.idle-anim {
    animation: breathing 1.8s ease-in-out infinite alternate;
}

/* 일하는 모드: 키보드 미친듯이 두드리기 */
.working-anim {
    animation: typing 0.15s ease-in-out infinite alternate;
}

@keyframes breathing {
    0% { transform: translateY(0px) scale(1); }
    100% { transform: translateY(-5px) scale(1.03); }
}

@keyframes typing {
    0% { transform: translateY(-2px) rotate(-3deg); }
    100% { transform: translateY(2px) rotate(3deg); }
}

.dialogue-box {
    flex-grow: 1;
    background: #ffffff;
    color: #222;
    padding: 12px 16px;
    border-radius: 8px;
    border: 3px solid #111;
    position: relative;
    box-shadow: 3px 3px 0px rgba(0,0,0,0.3);
}

.dialogue-box::before {
    content: '';
    position: absolute;
    left: -12px;
    top: 50%;
    transform: translateY(-50%);
    border-width: 6px 12px 6px 0;
    border-style: solid;
    border-color: transparent #111 transparent transparent;
}

.dialogue-name {
    font-size: 12px;
    font-weight: bold;
    color: #4a4a75;
    margin-bottom: 4px;
}

.dialogue-text {
    font-size: 14px;
    font-weight: 600;
    line-height: 1.4;
    margin: 0;
}
</style>
"""
st.markdown(pixel_art_css, unsafe_allow_html=True)

# 상단 타이틀
st.title("🇪🇺 유럽 랜드사 일정표 자동화 시스템")

# Gemini 클라이언트 초기화
api_key = st.secrets.get("GEMINI_API_KEY")
client = genai.Client(api_key=api_key) if api_key else None

col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1. 기본 정보 입력")
    start_date = st.date_input("출발일자 선택", datetime.today())
    tour_title = st.text_input("행사명", placeholder="예: 2026-12-05 SM투어) 북유럽 4국 10일 SK")
    
    routing_text = st.text_area(
        "루팅 및 일정 입력 (복사-붙여넣기)",
        height=320,
        placeholder="""1. 인천/코펜하겐 (SK988, 23:35출발)
2. 코펜하겐(관광, 뉘하운 유람선 탑승 필수)-DFDS탑승 - 오슬로 이동
3. 오슬로(관광, 시청사 내부 필수)-게일로
4. 게일로-베르겐(관광, 플뢰엔산 전망대 필수)
5. 베르겐-구드방겐(송네피요르드 유람선 탑승)-플람하선(플람열차 탑승)-플라로 이동
6. 플라-오슬로(관광, 뭉크미술관 필수) - 칼스타드
7. 칼스타드-스톡홀름(관광, 드로트닝홀름 궁전, 바사박물관 필수)
8. 스톡홀름(관광, 스톡홀름 시청사 및 왕실 무기고 필수) - 실자라인 탑승 - 헬싱키로 이동
9. 헬싱키(하선 후 관광) - 공항이동 - 헬싱키/코펜하겐/인천 (SK1717 / SK987 스케쥴)
10. 인천도착"""
    )

with col2:
    st.subheader("2. AI 막내 OP 작업실")
    
    # 2D 캐릭터 인터랙션 전용 컨테이너
    status_placeholder = st.empty()

    # 기본 대기 상태 연출 렌더링
    def render_assistant(status="idle", message="선배님, 루팅 던져주시면 바로 랜드사 규격 엑셀 뽑아오겠습니다! ☕"):
        anim_class = "working-anim" if status == "working" else "idle-anim"
        emoji = "👨‍💻" if status == "working" else "☕"
        status_label = "🔥 불타는 작업 중..." if status == "working" else "💤 대기 중"
        badge_bg = "#ff5e7e" if status == "working" else "#4caf50"
        
        status_html = f"""
        <div class="retro-container">
            <div class="retro-header">
                <span style="font-size: 13px; font-weight: bold;">🏢 OP 룸 • AI 막내 인턴 '김오피'</span>
                <span class="pixel-badge" style="background: {badge_bg};">{status_label}</span>
            </div>
            <div class="retro-stage">
                <div class="character-box {anim_class}">
                    <span>{emoji}</span>
                </div>
                <div class="dialogue-box">
                    <div class="dialogue-name">김오피 (유럽 수배 1년차 AI 인턴)</div>
                    <div class="dialogue-text">{message}</div>
                </div>
            </div>
        </div>
        """
        status_placeholder.markdown(status_html, unsafe_allow_html=True)

    render_assistant("idle")

    generate_btn = st.button("🚀 엑셀 일정표 생성하기", type="primary", use_container_width=True)

    if generate_btn:
        if not api_key:
            st.error("설정(Secrets)에서 GEMINI_API_KEY가 등록되지 않았습니다.")
        elif not routing_text.strip():
            st.warning("루팅 정보를 입력해 주세요.")
        else:
            # 1단계 캐릭터 반응: 작업 시작
            render_assistant("working", "선배님 확인했습니다! 루팅 분석 들어가겠습니다. 잠시만요!")
            time.sleep(0.7)
            
            try:
                render_assistant("working", "🛫 항공 블록 및 구간 이동 동선(페리/열차) 계산 중입니다...")
                
                prompt = f"""
당신은 15년 차 유럽 전문 랜드사의 수석 오퍼레이터입니다.
아래 루팅을 분석하여 한국 대형 여행사용 정밀 일정표 데이터를 JSON 배열로 출력하세요.

[입력된 루팅]
{routing_text}

[세부 정제 및 절대 규칙]
1. 도시명(city) 표기 규칙 (매우 중요):
   - 같은 도시의 일정이 계속 이어지는 동안에는 city 칸을 무조건 빈칸("")으로 둡니다.
   - 도시가 새롭게 바뀔 때만 그 행의 city 칸에 최초 1회만 적으세요.
2. 시간란(time) 표기 규칙 (매우 엄격):
   - '시간' 칸에는 오직 항공편 출발/도착 시간(HH:MM)과 관광 시작 행의 '전  일'만 기재합니다.
   - 페리 탑승, 열차, 버스 이동, 관광지 등에는 일체 시간을 적지 말고 무조건 빈칸("")으로 비워두세요.
   - '오전', '오후', '출발', '도착' 등의 한글 텍스트는 시간란에 절대 금지입니다.
3. 교통편('전용차량') 규칙:
   - 일반 육로 관광일의 경우 2행(조식 후 첫 관광/이동 시작 행)에만 transport="전용차량", time="전  일"을 단 1회 표기합니다.
   - 이후의 모든 행(관광지, 식사, 중간이동, 공항이동 등)의 transport는 무조건 빈칸("")입니다.
4. 일정 문구 작성 규칙:
   - '핵심관광', '집중투어' 같은 수식어 절대 금지. 오직 '[도시명] 시내 관광' 형태로 작성하세요 (예: "스톡홀름 시내 관광").
   - 필수 명소를 포함하여 해당 도시의 대표 명소(▶)를 최소 4개 이상 풍성하게 구성하세요.
   - 관광지는 반드시 '~한 [명소명]' 형태의 개별 행으로 분리 작성하세요.
5. 식사(meals) 규칙:
   - 특수 명칭(피요르드식, 선상식 등)은 쓰지 말고 기본적으로 "호텔식", "현지식"으로 통일하세요.
   - 기내식 또는 미제공만 "-"로 처리합니다. (예: ["호텔식", "현지식", "현지식"])
6. 숙박 표기:
   - 호텔: "HOTEL : ANY 4* HOTEL CLASS" (마지막 행="석식 후", 그 다음 행="호텔 투숙 및 휴식")
   - 크루즈/페리: "CRUISE : DFDS SEAWAYS (SEAVIEW 2-BERTH)" 등 (마지막 행="석식 후", 그 다음 행="선내 투숙 및 휴식")
   - 기내박: "HOTEL : 기내 숙박" (별도 투숙 문구 없음)
   - 귀국일: hotel="" (마지막 행="인천 도착 후 해산")

[반드시 아래 예시 포맷의 깔끔한 JSON 구조로 응답]
[
  {{
    "day": 7,
    "rows": [
      {{"city": "", "transport": "", "time": "", "schedule": "호텔 조식 후"}},
      {{"city": "스톡홀름", "transport": "전용차량", "time": "전  일", "schedule": "스톡홀름 시내 관광"}},
      {{"city": "", "transport": "", "time": "", "schedule": " ▶ 호화로운 스웨덴 왕실의 별궁 드로트닝홀름 궁전(내부)"}},
      {{"city": "", "transport": "", "time": "", "schedule": " ▶ 현존하는 가장 오래된 전함 바사호가 전시된 바사 박물관(내부)"}},
      {{"city": "", "transport": "", "time": "", "schedule": " ▶ 중세의 정취가 그대로 보존된 매력적인 감라스탄 구시가지"}},
      {{"city": "", "transport": "", "time": "", "schedule": " ▶ 웅장한 바로크 양식의 스톡홀름 왕궁(외관)"}},
      {{"city": "", "transport": "", "time": "", "schedule": "석식 후"}},
      {{"city": "", "transport": "", "time": "", "schedule": "호텔 투숙 및 휴식"}}
    ],
    "meals": ["호텔식", "현지식", "현지식"],
    "hotel": "HOTEL : ANY 4* HOTEL CLASS"
  }}
]
"""
                response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json"
                    )
                )
                
                render_assistant("working", "📊 랜드사 전용 엑셀 서식(18.75pt / 테두리 / 호텔 음영) 입히는 중입니다...")
                
                parsed_days = json.loads(response.text)
                
                # 엑셀 워크북 생성
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "일정표"
                ws.views.sheetView[0].showGridLines = True
                
                font_title = Font(name="맑은 고딕", size=15, bold=True)
                font_header = Font(name="맑은 고딕", size=10, bold=True)
                font_body = Font(name="맑은 고딕", size=9)
                font_bold = Font(name="맑은 고딕", size=9, bold=True)
                
                fill_title = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
                fill_hotel = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")
                
                align_center = Alignment(horizontal="center", vertical="center")
                align_mid_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
                thin_side = Side(style="thin", color="000000")
                med_side = Side(style="medium", color="000000")
                
                # 1행: 타이틀바
                full_title = tour_title.strip() if tour_title.strip() else f"{start_date.strftime('%Y-%m-%d')} 유럽투어 일정표"
                ws.row_dimensions[1].height = 28.0
                ws.cell(1, 1, full_title).font = font_title
                ws.cell(1, 1).alignment = align_center
                
                for c in range(1, 8):
                    cell = ws.cell(1, c)
                    cell.fill = fill_title
                    cell.border = Border(top=med_side, bottom=med_side, left=med_side if c==1 else None, right=med_side if c==7 else None)
                ws.merge_cells("A1:G1")
                
                # 2행: 간격행 (높이 13.0)
                ws.row_dimensions[2].height = 13.0
                
                # 3행: 헤더행
                headers = ["날짜", "도시", "교통편", "시간", "일정", "식사", ""]
                ws.row_dimensions[3].height = 22.0
                for c in range(1, 8):
                    cell = ws.cell(3, c, headers[c-1])
                    cell.font = font_header
                    cell.alignment = align_center
                    cell.border = Border(top=med_side, bottom=med_side, left=med_side if c==1 else thin_side, right=med_side if c==7 else thin_side)
                ws.merge_cells("F3:G3")
                
                # 열 너비 설정
                col_widths = {1: 9.25, 2: 12.75, 3: 7.5, 4: 9.25, 5: 64.25, 6: 3.63, 7: 10.0}
                for c_idx, width in col_widths.items():
                    ws.column_dimensions[get_column_letter(c_idx)].width = width
                
                days_korean = ["월요일", "화요일", "수요일", "목요일", "금요일", "토요일", "일요일"]
                cur_row = 4
                
                for item in parsed_days:
                    day_num = item.get("day", 1)
                    curr_date = start_date + timedelta(days=day_num - 1)
                    day_str = f"제{day_num:02d}일"
                    date_str = curr_date.strftime("%m월 %d일")
                    weekday_str = days_korean[curr_date.weekday()]
                    
                    rows_data = item.get("rows", [])
                    hotel_str = item.get("hotel", "")
                    meals = item.get("meals", ["-", "-", "-"])
                    
                    has_hotel = bool(hotel_str.strip())
                    total_rows = max(len(rows_data), 3) + (1 if has_hotel else 0)
                    
                    for r_idx in range(total_rows):
                        r = cur_row + r_idx
                        ws.row_dimensions[r].height = 18.75
                        is_last_row = (r_idx == total_rows - 1)
                        
                        for c in range(1, 8):
                            cell = ws.cell(r, c)
                            cell.font = font_body
                            
                            b_top = thin_side if (is_last_row and has_hotel) else None
                            b_bottom = med_side if is_last_row else None
                            b_left = med_side if c == 1 else thin_side
                            b_right = med_side if c == 7 else thin_side
                            
                            if is_last_row and has_hotel:
                                cell.fill = fill_hotel
                            cell.border = Border(top=b_top, bottom=b_bottom, left=b_left, right=b_right)
                    
                    # 1열 날짜
                    ws.cell(cur_row, 1, day_str).font = font_bold
                    ws.cell(cur_row, 1).alignment = align_center
                    ws.cell(cur_row + 1, 1, date_str).alignment = align_center
                    ws.cell(cur_row + 2, 1, weekday_str).alignment = align_center
                    
                    # 2~5열 세부 일정
                    for idx, r_data in enumerate(rows_data):
                        r = cur_row + idx
                        if r_data.get("city"):
                            cell_city = ws.cell(r, 2, r_data.get("city"))
                            cell_city.alignment = align_center
                        if r_data.get("transport"):
                            cell_tr = ws.cell(r, 3, r_data.get("transport"))
                            cell_tr.alignment = align_center
                        if r_data.get("time"):
                            cell_tm = ws.cell(r, 4, r_data.get("time"))
                            cell_tm.alignment = align_center
                        cell_sc = ws.cell(r, 5, r_data.get("schedule", ""))
                        cell_sc.alignment = align_mid_left
                    
                    # HOTEL 행
                    if has_hotel:
                        h_row = cur_row + total_rows - 1
                        h_cell = ws.cell(h_row, 5, hotel_str)
                        h_cell.font = font_bold
                        h_cell.alignment = align_mid_left
                    
                    # 식사 (조/중/석)
                    meal_labels = ["조", "중", "석"]
                    for m_idx in range(3):
                        ws.cell(cur_row + m_idx, 6, meal_labels[m_idx]).alignment = align_center
                        ws.cell(cur_row + m_idx, 7, meals[m_idx] if m_idx < len(meals) else "-").alignment = align_center
                    
                    cur_row += total_rows
                
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                
                # 완성 후 성공 대사로 전환
                render_assistant("idle", "선배님, 완벽하게 사내 규격으로 엑셀 파일 뽑아왔습니다! 아래에서 다운받아 검토해 주세요! 🎉")
                
                final_title = tour_title.strip() if tour_title.strip() else "유럽투어"
                file_name = f"{start_date.strftime('%Y-%m-%d')} {final_title} - 일정표.xlsx"
                
                st.download_button(
                    label="📥 완성된 사내 표준 엑셀 일정표 다운로드",
                    data=output,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True
                )
            except Exception as e:
                render_assistant("idle", f"헉 선배님... 작업 중에 오류가 발생했습니다 ㅠㅠ: {str(e)}")
                st.error(f"오류 내용: {str(e)}")
